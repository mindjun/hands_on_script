# read a docx content to a excel

import os
from docx import Document
from openpyxl import Workbook


def read_docs(path: str):
    document = Document(path)
    content = list()
    temp = list()
    for paragraph in document.paragraphs:
        if len(temp) == 0:
            text = ''.join(paragraph.text.split('.')[-1]).replace(' ', '')
        else:
            text = paragraph.text
        temp.append(text)
        if "【答案】" in text:
            if len(text) != 5:
                temp.insert(0, "多选题")
            else:
                temp.insert(0, "单选题")
            temp.insert(2, temp.pop(-2).replace('【解析】', ''))
            if len(temp) != 8:
                result = []
                temp_result = temp[3:]
                temp = temp[:3]
                for c in temp_result:
                    result.extend(c.split('\t'))
                result[-1] = temp_result[-1].replace('【答案】', '')
                temp.extend(result)
            else:
                temp[-1] = temp[-1].replace('【答案】', '')
            real_result = temp.pop()
            f_result = list()
            for a in temp:
                s = a.replace('A.', '').replace('B.', '').replace('C.', '').replace('D.', '').strip()
                f_result.append(s)
            f_result.insert(3, real_result)
            print(f_result)
            content.append(f_result)
            temp = list()
    print(content)
    return content


def to_excel(content: list):
    header = "题目类型,题目,解析,正确答案,答案A,答案B,答案C,答案D".split(",")
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for each_c in content:
        ws.append(each_c)

    res_path = os.path.join(cdw, 'template.xlsx')
    wb.save(res_path)


if __name__ == '__main__':
    cdw = os.getcwd()
    p = os.path.join(cdw, '2019_12.docx')
    word_c = read_docs(p)
    to_excel(word_c)
